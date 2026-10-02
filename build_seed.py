import json
import os
import sys

import openpyxl

DOWNLOADS = os.path.join(os.path.expanduser("~"), "Downloads")
EXCEL_CANDIDATOS = (
    "LOCALES_Y_MESAS_DISTRITO_CHULUCANAS_OCTUBRE_2026 - copia.xlsx",
    "LOCALES_Y_MESAS_DISTRITO_CHULUCANAS_OCTUBRE_2026.xlsx",
)
SHEET_MESAS = "data de mesas"

COLS = {
    "codigo_lv": 7,
    "nombre_local": 8,
    "direccion": 9,
    "orden_mesa": 12,
    "mesa_sufragio": 13,
    "aula": 14,
    "electores": 17,
    "personero": 20,
    "dni_personero": 21,
    "celular_personero": 22,
    "coordinador": 23,
}


def resolve_excel(path=None):
    if path:
        if not os.path.exists(path):
            raise FileNotFoundError(f"No se encontro el Excel: {path}")
        return path
    for nombre in EXCEL_CANDIDATOS:
        candidato = os.path.join(DOWNLOADS, nombre)
        if os.path.exists(candidato):
            return candidato
    raise FileNotFoundError(
        f"No se encontro ningun Excel en {DOWNLOADS}. "
        f"Esperado: {', '.join(EXCEL_CANDIDATOS)}"
    )


def texto(value):
    if not value:
        return ""
    return str(value).strip()


def numero_mesa(value):
    digits = "".join(ch for ch in texto(value) if ch.isdigit())
    return int(digits) if digits else 0


def dni_normalizado(value):
    digits = "".join(ch for ch in texto(value) if ch.isdigit())
    return digits.zfill(8) if digits else ""


def build(excel_path=None):
    excel_path = resolve_excel(excel_path)
    print(f">> Leyendo {excel_path}")

    wb = openpyxl.load_workbook(excel_path, read_only=True, data_only=True)
    if SHEET_MESAS not in wb.sheetnames:
        raise ValueError(
            f"La hoja '{SHEET_MESAS}' no existe. Hojas encontradas: {wb.sheetnames}"
        )
    ws = wb[SHEET_MESAS]

    locales = {}
    mesas = {}
    direcciones = {}
    coordinadores = {}
    personeros = []
    dnis_vistos = set()
    mesas_vistas = set()
    duplicados_eliminados = []

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not texto(row[COLS["nombre_local"]]):
            continue

        codigo = texto(row[COLS["codigo_lv"]])
        nombre_local = texto(row[COLS["nombre_local"]])
        mesa_num = numero_mesa(row[COLS["mesa_sufragio"]])
        if not codigo or not mesa_num:
            continue

        locales.setdefault(codigo, nombre_local)
        mesas.setdefault(codigo, [])
        direcciones.setdefault(codigo, texto(row[COLS["direccion"]]))

        clave_mesa = (codigo, mesa_num)
        if clave_mesa not in mesas_vistas:
            mesas_vistas.add(clave_mesa)
            mesas[codigo].append(mesa_num)

        coordinador = texto(row[COLS["coordinador"]])
        if coordinador and codigo not in coordinadores:
            coordinadores[codigo] = {
                "nombre": coordinador,
                "dni": dni_normalizado(row[COLS.get("dni_coordinador", 24)]),
                "celular": texto(row[COLS.get("celular_coordinador", 25)]),
            }

        nombre_personero = texto(row[COLS["personero"]])
        dni = dni_normalizado(row[COLS["dni_personero"]])
        if not nombre_personero or not dni:
            continue

        if dni in dnis_vistos:
            duplicados_eliminados.append(
                {
                    "dni": dni,
                    "personero": nombre_personero,
                    "mesa": mesa_num,
                    "local": nombre_local,
                }
            )
            continue

        dnis_vistos.add(dni)
        personeros.append(
            [
                codigo,
                mesa_num,
                dni,
                nombre_personero,
                texto(row[COLS["celular_personero"]]),
            ]
        )

    for codigo in mesas:
        mesas[codigo].sort()

    data = {
        "locales": locales,
        "direcciones": direcciones,
        "mesas": mesas,
        "personeros": personeros,
        "coordinadores": coordinadores,
    }

    print("== Resumen ==")
    print(f"Locales: {len(locales)}")
    print(f"Mesas: {sum(len(v) for v in mesas.values())}")
    print(f"Personeros: {len(personeros)}")
    print(f"Mesas sin personero: {sum(len(v) for v in mesas.values()) - len(personeros)}")
    print(f"Locales con coordinador: {len(coordinadores)}")
    print(f"Personeros duplicados eliminados: {len(duplicados_eliminados)}")
    for dup in duplicados_eliminados:
        print(
            f"   DNI {dup['dni']} | {dup['personero']} "
            f"| Mesa {dup['mesa']} @ {dup['local']}"
        )
    return data


def main():
    excel_path = sys.argv[1] if len(sys.argv) > 1 else None
    out_path = sys.argv[2] if len(sys.argv) > 2 else os.path.join(
        os.path.dirname(os.path.abspath(__file__)), "seed_data.json"
    )
    data = build(excel_path)
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"Escribi {out_path}")


if __name__ == "__main__":
    main()